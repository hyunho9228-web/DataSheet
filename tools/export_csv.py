"""구글 시트에서 받은 xlsx를 탭별 CSV로 내보낸다.

사용법:
    1) 구글 시트 > 파일 > 다운로드 > Microsoft Excel(.xlsx)
    2) python tools/export_csv.py 받은파일.xlsx
       -> csv/<탭이름>.csv 로 덮어씀. xlsx는 저장소에 넣지 않는다(.gitignore).

시트 규칙(README 참고): 1~5행은 설명, 6행 = 열 이름, 7행부터 데이터.
내보내면서 검사하는 것: A열 ID 중복, 빈 ID, 빈 열 이름.
"""
import csv
import sys
from pathlib import Path

from openpyxl import load_workbook

HEADER_ROW = 6
# 설명·계산용 탭은 내보내지 않음
SKIP_SHEETS = {"README", "Reference", "REGISTRY", "StatCurve", "Balance"}

OUT_DIR = Path(__file__).resolve().parent.parent / "csv"


def cell_text(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def export_sheet(ws):
    rows = list(ws.iter_rows(min_row=HEADER_ROW, values_only=True))
    if not rows:
        return None, ["데이터 없음"]
    header = [cell_text(v) for v in rows[0]]
    while header and header[-1] == "":
        header.pop()
    width = len(header)
    errors = []
    if "" in header:
        errors.append(f"6행에 빈 열 이름이 있음 (열 {header.index('') + 1})")

    data, seen = [], {}
    for i, r in enumerate(rows[1:], start=HEADER_ROW + 1):
        vals = [cell_text(v) for v in r[:width]]
        vals += [""] * (width - len(vals))
        if not any(vals):
            continue
        rid = vals[0]
        if rid == "":
            errors.append(f"{i}행: ID(A열)가 비어 있음")
        elif rid in seen:
            errors.append(f"{i}행: ID {rid} 중복 ({seen[rid]}행과 같음)")
        else:
            seen[rid] = i
        data.append(vals)
    return [header] + data, errors


def main():
    if len(sys.argv) != 2:
        print(__doc__)
        sys.exit(1)
    wb = load_workbook(sys.argv[1], data_only=True, read_only=True)
    OUT_DIR.mkdir(exist_ok=True)
    failed = False
    for ws in wb.worksheets:
        if ws.title in SKIP_SHEETS:
            continue
        table, errors = export_sheet(ws)
        for e in errors:
            print(f"[오류] {ws.title}: {e}")
        if errors:
            failed = True
            continue
        path = OUT_DIR / f"{ws.title}.csv"
        with open(path, "w", encoding="utf-8", newline="") as f:
            csv.writer(f, lineterminator="\n").writerows(table)
        print(f"{ws.title}: {len(table) - 1}행 -> {path.relative_to(OUT_DIR.parent)}")
    if failed:
        print("오류가 있는 탭은 내보내지 않았습니다. 시트를 고친 뒤 다시 실행하세요.")
        sys.exit(1)


if __name__ == "__main__":
    main()
